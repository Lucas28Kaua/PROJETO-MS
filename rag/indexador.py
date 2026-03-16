import os
import fitz  # pymupdf
import openpyxl
import docx
import io
import chromadb
from sentence_transformers import SentenceTransformer

CHROMA_PATH = os.path.join(os.path.abspath(os.path.dirname(__file__)), '..', 'chroma_db')

client_chroma = chromadb.PersistentClient(path=CHROMA_PATH)
colecao = client_chroma.get_or_create_collection(name="mscred_docs")
modelo_embedding = SentenceTransformer('all-MiniLM-L6-v2')

def extrair_texto(arquivo_bytes, nome):
    nome = nome.lower()

    if nome.endswith(".pdf"):
        doc = fitz.open(stream=arquivo_bytes, filetype="pdf")
        return "\n".join(page.get_text() for page in doc)

    elif nome.endswith((".xlsx", ".xls")):
        wb = openpyxl.load_workbook(io.BytesIO(arquivo_bytes), data_only=True)
        texto = []
        for sheet in wb.worksheets:
            texto.append(f"[Aba: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                linha = " | ".join(str(c) for c in row if c is not None)
                if linha:
                    texto.append(linha)
        return "\n".join(texto)

    elif nome.endswith(".docx"):
        doc = docx.Document(io.BytesIO(arquivo_bytes))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    elif nome.endswith(".txt"):
        return arquivo_bytes.decode("utf-8")

    return ""

def chunkar_texto(texto, tamanho=500, sobreposicao=50):
    palavras = texto.split()
    chunks = []
    i = 0
    while i <len(palavras):
        chunk = " ".join(palavras[i:i+tamanho])
        chunks.append(chunk)
        i += tamanho - sobreposicao
    return chunks

def indexar_documento(arquivo_bytes, nome_arquivo):
    remover_documento(nome_arquivo)

    texto = extrair_texto(arquivo_bytes, nome_arquivo)
    if not texto.strip():
        return {"erro": "Não foi possível extrair texto do documento"}

    chunks = chunkar_texto(texto)
    embeddings = modelo_embedding.encode(chunks).tolist()

    ids = [f"{nome_arquivo}__chunk_{i}" for i in range(len(chunks))]

    metadatas = [{"fonte": nome_arquivo, "chunk": i} for i in range(len(chunks))]

    colecao.add(documents=chunks, embeddings=embeddings, ids=ids, metadatas=metadatas)

    return {"sucesso": True, "chunks": len(chunks), "arquivo": nome_arquivo}

def remover_documento(nome_arquivo):
    try:
        existentes = colecao.get(where={"fonte":nome_arquivo})
        if existentes["ids"]:
            colecao.delete(ids=existentes['ids'])
    except:
        pass

def listar_documentos():
    try:
        todos = colecao.get()
        fontes = list(set(m["fonte"] for m in todos["metadatas"])) if todos["metadatas"] else []
        return sorted(fontes)
    except:
        return []